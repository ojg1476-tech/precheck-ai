"""일괄 점검 결과 → 엑셀 리포트 (요약 · 위험 순위 · 상세 · 기준 안내)."""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .batch import ProductResult, overview, penalty_text
from .engine import Checker

TYPE_NAME = {"cos": "일반 화장품", "fcos": "기능성 화장품", "food": "일반 식품", "hfood": "건강기능식품", "": "-"}
TIER_FILL = {"높음": "FDE8E6", "중간": "FDF0D9", "확인": "E6EEFB", "양호": "E2F3EA", "오류": "EEEEEE"}
GRADE_FILL = {"red": "FDE8E6", "amber": "FDF0D9", "blue": "E6EEFB"}
HEAD = PatternFill("solid", fgColor="1F5C4A")
THIN = Side(style="thin", color="DDE2DF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
NOTICE = ("이 리포트는 공개된 법령·식약처 기준·실제 단속 사례에 비춰 문제가 될 수 있는 표현을 알려 주는 점검 자료예요. "
          "법적으로 위반인지 확정하는 판단은 아니며, 최종 확인은 관계 기관이나 전문가에게 받으세요.")


def _table(ws, header: list[str], rows: list[list], widths: list[int], start_row: int = 1):
    for c, h in enumerate(header, 1):
        cell = ws.cell(row=start_row, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BOX
    for r, row in enumerate(rows, start_row + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = WRAP
            cell.border = BOX
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def build_report(results: list[ProductResult], checker: Checker | None = None, shop: str = "") -> bytes:
    ck = checker or Checker()
    ov = overview(results, ck)
    wb = Workbook()

    # 1. 요약
    ws = wb.active
    ws.title = "요약"
    ws["A1"] = f"상세페이지 광고 일괄 점검 리포트{(' — ' + shop) if shop else ''}"
    ws["A1"].font = Font(bold=True, size=15)
    ws["A2"] = f"점검일 {date.today().isoformat()} · 기준 버전 {ck.version} · 규칙 {len(ck.rules)}개"
    ws["A2"].font = Font(color="5D6B66")
    lines = [("점검한 상품", ov["total"])] + [(f"위험도 {k}", v) for k, v in ov["tiers"].items()]
    for i, (k, v) in enumerate(lines, 4):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)
        if k.startswith("위험도 "):
            ws.cell(row=i, column=1).fill = PatternFill("solid", fgColor=TIER_FILL[k[4:]])
    r = 4 + len(lines) + 1
    ws.cell(row=r, column=1, value="여러 상품에서 반복된 문제").font = Font(bold=True)
    for i, (name, n) in enumerate(ov["common"], r + 1):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=f"{n}개 상품")
    r = r + 2 + max(len(ov["common"]), 1)
    ws.cell(row=r, column=1, value="위험도 뜻").font = Font(bold=True)
    for i, t in enumerate(["높음: 꼭 고쳐야 할 표현이 하나 이상 있어요.", "중간: 문제가 될 수 있어 확인이 필요한 표현이 있어요.",
                           "확인: 쓸 수는 있지만 증명 자료나 인정 여부를 확인해야 하는 표현만 있어요.",
                           "양호: 기본 검사에서 걸린 표현이 없어요.", "같은 위험도 안에서는 형사처벌이 무거운 상품이 먼저 와요."], r + 1):
        ws.cell(row=i, column=1, value=t)
    ws.cell(row=r + 7, column=1, value=NOTICE).alignment = WRAP
    ws.merge_cells(start_row=r + 7, start_column=1, end_row=r + 7, end_column=4)
    ws.row_dimensions[r + 7].height = 48
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 14

    # 2. 위험 순위
    ws = wb.create_sheet("위험 순위")
    rows = []
    for i, p in enumerate(results, 1):
        c = p.counts
        rows.append([i, p.name, TYPE_NAME.get(p.product_type, p.product_type), p.tier, penalty_text(p.max_penalty),
                     c["red"], c["amber"], c["blue"], p.score, p.error or ""])
    _table(ws, ["순위", "상품명", "제품 종류", "위험도", "가장 무거운 형사처벌(최대)", "꼭 고치기", "확인 필요", "자료 챙기기", "점수", "메모"],
           rows, [6, 26, 14, 9, 30, 9, 9, 10, 7, 28])
    for i, p in enumerate(results, 2):
        ws.cell(row=i, column=4).fill = PatternFill("solid", fgColor=TIER_FILL[p.tier])

    # 3. 상세
    ws = wb.create_sheet("상세")
    rows, fills = [], []
    for p in results:
        for f in sorted(p.findings, key=lambda x: ({"red": 0, "amber": 1, "blue": 2}[x.grade], x.start)):
            rows.append([p.name, f.grade_name, f.phrase, f.label, f.why, f.alt,
                         "\n".join("□ " + c for c in f.checks), f.penalty, f.law, f.source])
            fills.append(f.grade)
    _table(ws, ["상품명", "등급", "찾은 표현", "쉬운 이름", "왜 문제인가", "대신 이렇게", "쓰기 전에 확인할 것", "이대로 두면", "관련 기준", "출처"],
           rows, [20, 10, 18, 20, 36, 26, 44, 44, 30, 28])
    for i, g in enumerate(fills, 2):
        ws.cell(row=i, column=2).fill = PatternFill("solid", fgColor=GRADE_FILL[g])

    # 4. 기준 안내
    ws = wb.create_sheet("기준 안내")
    ws["A1"] = "이 점검의 기준과 출처"
    ws["A1"].font = Font(bold=True, size=13)
    basis = [
        "화장품: 화장품법 제13조·제14조·제37조, 시행규칙 별표5(표시·광고 준수사항)·별표7(행정처분 기준), 대한화장품협회 「화장품 광고자문 기준 및 해설서」(2025)",
        "식품: 식품 등의 표시·광고에 관한 법률 제8조·제26조·제27조, 시행령 별표1(부당한 표시·광고의 내용), 시행규칙 별표7(행정처분 기준)",
        "형사처벌은 법에 정해진 최대치예요. 실제 처벌은 사안마다 달라요.",
        "행정처분 기간은 처음 위반했을 때 기준이고, 1년 안에 반복하면 늘어나요. 식품은 판매업 기준이에요.",
        "‘꼭 고치기’는 기준에 금지로 적힌 표현, ‘확인 필요’는 맥락에 따라 위반이 될 수 있는 표현, ‘자료 챙기기’는 증명 자료가 있어야 쓰는 표현이에요.",
        NOTICE,
    ]
    for i, t in enumerate(basis, 3):
        ws.cell(row=i, column=1, value=t).alignment = WRAP
        ws.row_dimensions[i].height = 36
    ws.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
