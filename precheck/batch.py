"""여러 상품 일괄 점검: 상품 목록 → 상품별 검사 결과 → 위험 순위.

입력 표의 열 이름은 한국어/영어 몇 가지를 받아준다.
  상품명(name) · 제품종류(type) · 상세문구(text)
제품종류는 '일반 화장품', '기능성 화장품', '일반 식품', '건강기능식품' 또는 cos/fcos/food/hfood.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path

from .engine import Checker, Finding, summarize

TYPE_ALIASES = {
    "cos": "cos", "일반화장품": "cos", "화장품": "cos",
    "fcos": "fcos", "기능성화장품": "fcos",
    "food": "food", "일반식품": "food", "식품": "food",
    "hfood": "hfood", "건강기능식품": "hfood", "건기식": "hfood",
}
COLUMN_ALIASES = {
    "name": ["상품명", "제품명", "name", "product", "상품"],
    "type": ["제품종류", "제품 종류", "종류", "type", "product_type", "구분"],
    "text": ["상세문구", "상세페이지 문구", "상세페이지", "문구", "text", "본문", "내용"],
}

# 형사처벌 최대치 (처분 키 → (징역 연수, 벌금 원)). f34는 판단에 따라 둘 중 하나라 무거운 쪽으로 본다.
PEN_MAX = {
    "cA": (1, 10_000_000), "cB": (1, 10_000_000), "cAB": (1, 10_000_000), "cP": (1, 10_000_000), "cX": (1, 10_000_000),
    "f1": (10, 100_000_000), "f2": (10, 100_000_000), "f3": (10, 100_000_000), "f34": (10, 100_000_000),
    "f4": (5, 50_000_000), "f5": (5, 50_000_000), "f6": (5, 50_000_000), "fP": (5, 50_000_000),
}
WEIGHT = {"red": 10, "amber": 4, "blue": 1}
TIERS = [("높음", "red"), ("중간", "amber"), ("확인", "blue")]


def normalize_type(value: str) -> str:
    key = str(value or "").replace(" ", "").strip().lower()
    if key not in TYPE_ALIASES:
        raise ValueError(f"제품 종류 '{value}'를 알 수 없어요. 일반 화장품·기능성 화장품·일반 식품·건강기능식품 중 하나로 적어 주세요.")
    return TYPE_ALIASES[key]


@dataclass
class ProductResult:
    row: int                      # 원본 표의 행 번호 (머리글 다음이 1)
    name: str
    product_type: str
    text: str
    findings: list[Finding] = field(default_factory=list)
    error: str = ""

    @property
    def counts(self) -> dict:
        return summarize(self.findings)

    @property
    def score(self) -> int:
        return sum(WEIGHT[f.grade] for f in self.findings)

    @property
    def tier(self) -> str:
        if self.error:
            return "오류"
        c = self.counts
        for name, grade in TIERS:
            if c[grade]:
                return name
        return "양호"

    @property
    def max_penalty(self) -> tuple[int, int] | None:
        """꼭 고치기·확인 필요 표현 중 가장 무거운 형사처벌 최대치."""
        best = None
        for f in self.findings:
            if f.grade == "blue":
                continue
            p = PEN_MAX.get(f.penalty_key)
            if p and (best is None or p > best):
                best = p
        return best

    def sort_key(self):
        order = {"높음": 0, "중간": 1, "확인": 2, "양호": 3, "오류": 4}
        p = self.max_penalty or (0, 0)
        return (order[self.tier], -p[0], -self.score, self.row)


def _pick(header: list[str], key: str) -> int | None:
    norm = [h.replace(" ", "").strip().lower() for h in header]
    for alias in COLUMN_ALIASES[key]:
        a = alias.replace(" ", "").lower()
        if a in norm:
            return norm.index(a)
    return None


def read_table(data: bytes, filename: str) -> list[dict]:
    """CSV 또는 엑셀 파일 내용을 [{row, name, type, text}] 로 읽는다."""
    suffix = Path(filename).suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = [[("" if c is None else str(c)) for c in r] for r in ws.iter_rows(values_only=True)]
    elif suffix in (".csv", ".txt"):
        text = None
        for enc in ("utf-8-sig", "cp949"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise ValueError("파일 글자 인코딩을 읽지 못했어요. UTF-8 CSV로 저장해 주세요.")
        rows = list(csv.reader(io.StringIO(text)))
    else:
        raise ValueError("CSV 또는 엑셀(.xlsx) 파일만 올릴 수 있어요.")
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise ValueError("파일이 비어 있어요.")
    header = [str(h) for h in rows[0]]
    idx = {k: _pick(header, k) for k in COLUMN_ALIASES}
    missing = [COLUMN_ALIASES[k][0] for k, v in idx.items() if v is None]
    if missing:
        raise ValueError(f"'{', '.join(missing)}' 열을 찾지 못했어요. 첫 줄에 상품명·제품종류·상세문구 열 이름을 적어 주세요.")
    out = []
    for i, r in enumerate(rows[1:], 1):
        get = lambda k: (r[idx[k]] if idx[k] < len(r) else "").strip()
        out.append({"row": i, "name": get("name") or f"상품 {i}", "type": get("type"), "text": get("text")})
    return out


def check_products(items: list[dict], checker: Checker | None = None) -> list[ProductResult]:
    ck = checker or Checker()
    results = []
    for it in items:
        pr = ProductResult(row=it["row"], name=it["name"], product_type="", text=it.get("text", ""))
        try:
            pr.product_type = normalize_type(it.get("type", ""))
            pr.findings = ck.check(pr.text, pr.product_type, name=pr.name) if pr.text else []
            if it.get("_err"):
                pr.error = f"이미지 읽기 실패: {it['_err'][:80]}"
            elif not pr.text:
                pr.error = "상세 문구가 비어 있어요."
        except ValueError as e:
            pr.error = str(e)
        results.append(pr)
    return sorted(results, key=ProductResult.sort_key)


def overview(results: list[ProductResult], checker: Checker | None = None) -> dict:
    cats = (checker or Checker()).categories
    tiers = {t: 0 for t in ("높음", "중간", "확인", "양호", "오류")}
    for r in results:
        tiers[r.tier] += 1
    phrases = {}
    for r in results:
        for f in r.findings:
            if f.grade != "blue":
                phrases.setdefault(cats.get(f.category, {}).get("name", f.label), set()).add(r.name)
    common = sorted(((k, len(v)) for k, v in phrases.items()), key=lambda x: -x[1])[:5]
    return {"total": len(results), "tiers": tiers, "common": common}


def penalty_text(p: tuple[int, int] | None) -> str:
    if not p:
        return "-"
    years, won = p
    money = f"{won // 100_000_000}억 원" if won >= 100_000_000 else f"{won // 10_000:,}만 원"
    return f"징역 {years}년 이하 또는 벌금 {money} 이하"
