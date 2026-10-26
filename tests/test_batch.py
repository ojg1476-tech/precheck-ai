"""일괄 점검·리포트·이미지 읽기 모듈 테스트 (실제 API 호출 없음). python tests/test_batch.py 로도 실행."""
import io
import json
import sys
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from PIL import Image  # noqa: E402

from precheck import check_products, read_table  # noqa: E402
from precheck import ocr  # noqa: E402
from precheck.batch import normalize_type, penalty_text  # noqa: E402
from precheck.report import build_report  # noqa: E402

SAMPLE = ROOT / "examples" / "sample_products.csv"


def test_read_and_rank():
    res = check_products(read_table(SAMPLE.read_bytes(), "s.csv"))
    tiers = [r.tier for r in res]
    assert tiers[0] == "높음" and tiers[-1] == "오류"          # 위험한 상품이 먼저, 빈 문구는 맨 끝
    assert tiers.index("양호") > tiers.index("확인")
    top = res[0]
    assert top.max_penalty == (10, 100_000_000)               # 식품 질병·의약품·건기식 오인 → 10년/1억
    names = {r.name: r.tier for r in res}
    assert names["현미 누룽지"] == "양호" and names["클렌징 폼"] == "양호"


def test_type_aliases():
    assert normalize_type("건강기능식품") == "hfood" and normalize_type(" 일반 화장품 ") == "cos"
    try:
        normalize_type("의약품")
        assert False
    except ValueError:
        pass


def test_xlsx_input_and_report():
    from openpyxl import Workbook, load_workbook
    wb = Workbook()
    wb.active.append(["제품명", "종류", "내용"])
    wb.active.append(["차", "식품", "감기 예방 차"])
    buf = io.BytesIO()
    wb.save(buf)
    res = check_products(read_table(buf.getvalue(), "a.xlsx"))
    assert res[0].tier == "높음"
    out = load_workbook(io.BytesIO(build_report(res)))
    assert out.sheetnames == ["요약", "위험 순위", "상세", "기준 안내"]
    assert out["위험 순위"].cell(row=2, column=2).value == "차"


def test_penalty_text():
    assert penalty_text((10, 100_000_000)) == "징역 10년 이하 또는 벌금 1억 원 이하"
    assert penalty_text((1, 10_000_000)) == "징역 1년 이하 또는 벌금 1,000만 원 이하"


def _png(w, h):
    b = io.BytesIO()
    Image.new("RGB", (w, h), "white").save(b, format="PNG")
    return b.getvalue()


def test_slice_long_image():
    chunks = ocr.slice_image(_png(860, 10000))
    sizes = [Image.open(io.BytesIO(c)).size for c in chunks]
    assert all(w == 860 and h <= ocr.CHUNK_HEIGHT for w, h in sizes)
    assert sum(h for _, h in sizes) >= 10000                   # 겹침 때문에 원본보다 같거나 길다
    assert Image.open(io.BytesIO(ocr.slice_image(_png(2000, 500))[0])).size[0] == ocr.MAX_WIDTH


def test_read_text_mocked():
    fake = {"content": [{"type": "text", "text": "감기 예방 차"}], "usage": {"input_tokens": 12000, "output_tokens": 300}}

    class R:
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def read(self): return json.dumps(fake).encode()

    with mock.patch("urllib.request.urlopen", return_value=R()) as m:
        r = ocr.read_text([_png(860, 4000)], api_key="test")
    sent = json.loads(m.call_args[0][0].data)
    assert sent["model"] == ocr.DEFAULT_MODEL and sum(c["type"] == "image" for c in sent["messages"][0]["content"]) == 3
    assert r.text == "감기 예방 차" and round(r.usd, 5) == round(12000 / 1e6 * 0.10 + 300 / 1e6 * 0.50, 5)


if __name__ == "__main__":
    fns = [v for k, v in dict(globals()).items() if k.startswith("test_")]
    for f in fns:
        f()
        print("통과:", f.__name__)
    print(f"{len(fns)}/{len(fns)} 통과")
